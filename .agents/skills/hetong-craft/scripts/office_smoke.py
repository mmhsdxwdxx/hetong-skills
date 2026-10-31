#!/usr/bin/env python3
"""Create and verify synthetic Word/redline/report samples, then render them locally."""
from __future__ import annotations
import argparse
import hashlib
import json
import os
import subprocess
import sys
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

SCRIPTS = Path(__file__).resolve().parent
W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'


def command(script: Path, *args: str) -> dict:
    env = dict(os.environ, PYTHONUTF8='1', PYTHONIOENCODING='utf-8')
    run = subprocess.run([sys.executable, str(script), *map(str, args)],
                         capture_output=True, text=True, encoding='utf-8', env=env)
    if run.returncode:
        raise RuntimeError(f'{script.name}: exit {run.returncode}: {(run.stderr + run.stdout)[-3000:]}')
    return json.loads(run.stdout)


def pdf_fonts(path: Path) -> set[str]:
    from pypdf import PdfReader
    result = set()
    for page in PdfReader(path).pages:
        resources = page.get('/Resources', {}).get_object()
        fonts = resources.get('/Font', {}).get_object()
        for reference in fonts.values():
            font = reference.get_object()
            if font.get('/BaseFont'):
                result.add(str(font['/BaseFont']).lstrip('/'))
    return result


def run(out: Path, require_fonts: bool = False) -> dict:
    from docx import Document
    from docx.shared import Mm, Pt
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    if out.exists() and any(out.iterdir()):
        raise ValueError('样张目录非空，请使用新目录')
    out.mkdir(parents=True, exist_ok=True)
    original = out / '虚构环境样张_原文.docx'
    revised = out / '虚构环境样张_修订批注.docx'
    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Mm(210), Mm(297)
    section.left_margin = section.right_margin = Mm(25.4)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = '宋体', Pt(10.5)
    normal.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '宋体')
    doc.add_heading('虚构Word环境样张', 0)
    target = '本文件用于环境检查，不是实际合同。'
    doc.add_paragraph(target)
    table = doc.add_table(rows=3, cols=4)
    table.autofit = False
    widths = [22, 52, 50, 35.2]
    values = [
        ['日期', '我方现金流及付款义务', '交付与担保', '发票（万元）'],
        ['T0', '当期应付：见示例条件\n实际已付：待核', '虚构交付节点\n虚构保函检查', '未设置开票事件'],
        ['T1', '退款分支单列\n不重复计算占用', '虚构验收节点', '实际开票情况待核'],
    ]
    for col, width in zip(table.columns, widths):
        col.width = Mm(width)
    for row, texts in zip(table.rows, values):
        for cell, width, text in zip(row.cells, widths, texts):
            cell.width, cell.text = Mm(width), text
    header = OxmlElement('w:tblHeader')
    table.rows[0]._tr.get_or_add_trPr().append(header)
    doc.save(original)
    before = hashlib.sha256(original.read_bytes()).hexdigest()
    redline = SCRIPTS / 'hetong_redline.py'
    inventory = command(redline, 'inspect', original)
    paragraph = next(p['id'] for p in inventory['paragraphs'] if p['text'] == target)
    edits = out / 'edits.json'
    edits.write_text(json.dumps([{'id': 'ENV-01', 'para': paragraph, 'old': '环境检查',
                                 'new': '修订和批注检查', 'reason': '虚构样张：核对原生修订及批注显示。'}], ensure_ascii=False), encoding='utf-8')
    plan = out / 'plan.json'
    command(redline, 'plan', original, edits, '-o', plan)
    command(redline, 'apply', original, plan, '-o', revised, '--author', '环境检查')
    verify = command(redline, 'verify', original, revised, plan)
    if not verify.get('ok'):
        raise RuntimeError('原生修订核验未通过')
    with zipfile.ZipFile(revised) as archive:
        body = ET.fromstring(archive.read('word/document.xml'))
        comments = ET.fromstring(archive.read('word/comments.xml'))
    counts = {'insertions': len(body.findall('.//' + W + 'ins')),
              'deletions': len(body.findall('.//' + W + 'del')),
              'comments': len(comments.findall(W + 'comment'))}
    if not all(counts.values()) or hashlib.sha256(original.read_bytes()).hexdigest() != before:
        raise RuntimeError('原生修订/批注或原件保留检查失败')
    report_root = SCRIPTS.parents[1] / 'hetong-fast-xqy-v1'
    data = json.loads((report_root / 'assets/example-report.json').read_text(encoding='utf-8'))
    data['title'] = '虚构XQY环境检查报告'
    data_path = out / '虚构报告数据.json'
    data_path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    report = out / '虚构XQY报告.docx'
    command(report_root / 'scripts/build_report.py', '--input', data_path, '--out', report)
    rendered = command(SCRIPTS / 'office_kit.py', 'render', original, revised, report,
                       '--out', out / '预览', '--dpi', '110')
    fonts = set()
    report_fonts = set()
    for entry in rendered['entries']:
        current = pdf_fonts(Path(entry['pdf']))
        fonts.update(current)
        if Path(entry['input']).resolve() == report.resolve():
            report_fonts.update(current)
    folded = ' '.join(report_fonts).lower().replace('-', '')
    exact = {'宋体': 'simsun' in folded or '宋体' in folded,
             '黑体': 'simhei' in folded or '黑体' in folded}
    result = {'ok': not require_fonts or all(exact.values()), 'synthetic_only': True,
              'source_unchanged': True, 'redline_verified': bool(verify.get('ok')),
              'native_marks': counts, 'rendered_documents': len(rendered['entries']),
              'pdf_font_names': sorted(fonts), 'required_report_fonts': exact,
              'report_pdf_font_names': sorted(report_fonts),
              'artifacts': str(out), 'native_word_wps_roundtrip': 'not performed'}
    if require_fonts and not all(exact.values()):
        result['error'] = '样张未实际使用宋体/黑体。请安装有权使用的相应字体；Windows可检查简体中文补充字体，然后重跑。'
    (out / 'smoke-result.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    return result


def main() -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, OSError):
            pass
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--out', required=True, type=Path)
    parser.add_argument('--require-fonts', action='store_true', help='宋体和黑体未实际用于PDF则失败')
    args = parser.parse_args()
    try:
        result = run(args.out.resolve(), args.require_fonts)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0 if result['ok'] else 1
    except Exception as exc:
        print(json.dumps({'ok': False, 'error': str(exc)}, ensure_ascii=False))
        return 1


if __name__ == '__main__':
    raise SystemExit(main())
