#!/usr/bin/env python3
"""Build the editable XQY review report; performs no contract or legal analysis."""
from __future__ import annotations

import argparse
import hashlib
import io
import json
import re
import sys
from datetime import date
from pathlib import Path
from typing import Any
from urllib.parse import urlparse

SCHEMA = 'hetong-fast-xqy-v1/report-1'
HEADINGS = {
    1: '基本情况', 2: '发票', 3: '履约保函', 4: '产品配套资料',
    5: '验收', 6: '违约责任', 7: '不可抗力', 8: '廉洁条款',
}
STATUSES = {'已具备', '部分具备', '缺失', '待核实', '不适用'}
FIXED_STATUSES = {'符合', '部分符合', '缺失', '冲突', '待核实', '本次明确偏离'}
ISSUE_TYPES = {'revise': '建议修订', 'attention': '履约注意', 'verify': '待核实'}
CATALOG_DIR = Path(__file__).resolve().parents[2] / 'hetong-craft' / 'references'


class ReportError(ValueError):
    pass


def required_text(obj: dict, key: str, where: str) -> str:
    value = obj.get(key)
    if not isinstance(value, str) or not value.strip():
        raise ReportError(f'{where}.{key}须为非空文本')
    if any(ord(c) < 32 and c not in '\n\t' for c in value):
        raise ReportError(f'{where}.{key}含非法控制字符')
    return value.strip()


def optional_texts(obj: dict, keys: tuple[str, ...], where: str) -> None:
    for key in keys:
        if key in obj and obj[key] != '':
            required_text(obj, key, where)


def load_catalog(role: str) -> dict:
    if role not in ('buyer', 'seller'):
        raise ReportError('stance.role须为buyer或seller')
    path = CATALOG_DIR / f'{role}-clause-catalog.json'
    try:
        return json.loads(path.read_text(encoding='utf-8'))
    except (OSError, ValueError) as exc:
        raise ReportError(f'本方固定条款索引不可读：{path}') from exc


def report_items(data: dict) -> list[dict]:
    return [data['basics'], *data['sections'], *data.get('extras', []), *data.get('issues', [])]


def validate_issues(data: dict) -> None:
    issues = data.get('issues', [])
    if not isinstance(issues, list):
        raise ReportError('issues须为数组')
    seen = set()
    for issue in issues:
        if not isinstance(issue, dict):
            raise ReportError('issues每项须为对象')
        code = required_text(issue, 'id', 'issue')
        if not re.fullmatch(r'R-\d{3,}', code) or int(code[2:]) == 0 or code in seen:
            raise ReportError('事项编号须为不重复的R-001格式，禁止重编或重复')
        seen.add(code)
        section = issue.get('section')
        if not ((type(section) is int and 1 <= section <= 8) or section == 'extra'):
            raise ReportError(f'{code}.section须为1—8或extra')
        if issue.get('type') not in ISSUE_TYPES:
            raise ReportError(f'{code}.type须为revise、attention或verify')
        for key in ('title', 'location', 'detail', 'suggestion', 'comment_summary'):
            required_text(issue, key, code)
        optional_texts(issue, ('original_excerpt',), code)
        anchors = issue.get('anchors', [])
        if not isinstance(anchors, list):
            raise ReportError(f'{code}.anchors须为数组')
        for anchor in anchors:
            if not isinstance(anchor, dict) or not re.fullmatch(r'RL\d{4,}', str(anchor.get('para', ''))):
                raise ReportError(f'{code}批注锚点须使用实际inspect的RL编号')
            required_text(anchor, 'text', code + '.anchor')


def fixed_references(item: dict, data: dict) -> list[tuple[dict, dict]]:
    matches = item.get('fixed_clauses', [])
    if not isinstance(matches, list):
        raise ReportError('fixed_clauses须为数组')
    if not matches:
        return []
    stance = data.get('stance')
    if not isinstance(stance, dict):
        raise ReportError('标记固定条款前必须明确stance，不能猜测本方')
    catalog = load_catalog(stance.get('role'))
    modules = {m['id']: m for m in catalog['modules']}
    seen = set()
    result = []
    for match in matches:
        if not isinstance(match, dict):
            raise ReportError('fixed_clauses每项须为对象')
        code = required_text(match, 'id', 'fixed_clause')
        if code not in modules:
            raise ReportError(f'固定条款{code}不属于本方当前条款库，禁止跨方或虚构编号')
        if code in seen:
            raise ReportError(f'同一要点中固定条款{code}重复')
        seen.add(code)
        if match.get('status') not in FIXED_STATUSES:
            raise ReportError(f'{code}须明确固定条款符合情况')
        required_text(match, 'analysis', code)
        module = modules[code]
        if 'categories' in stance and module['category'] not in stance['categories']:
            raise ReportError(f'{code}超出本次明确的固定条款类别范围')
        result.append((match, module))
    return result


def validate(data: Any) -> dict:
    if not isinstance(data, dict) or data.get('schema') != SCHEMA:
        raise ReportError(f'schema必须为{SCHEMA}')
    for key in ('title', 'review_date', 'perspective'):
        required_text(data, key, 'report')
    try:
        if date.fromisoformat(data['review_date']).isoformat() != data['review_date']:
            raise ValueError()
    except ValueError as exc:
        raise ReportError('review_date须为有效YYYY-MM-DD') from exc
    basics = data.get('basics')
    if not isinstance(basics, dict):
        raise ReportError('basics须为对象')
    for key in ('product', 'quantity', 'price', 'delivery'):
        required_text(basics, key, 'basics')
    optional_texts(basics, ('conclusion',), 'basics')
    sections = data.get('sections')
    if not isinstance(sections, list) or len(sections) != 7:
        raise ReportError('sections须完整包含第2至8项，不以省略代替不适用')
    if any(not isinstance(s, dict) for s in sections):
        raise ReportError('sections每项须为对象')
    if [s.get('id') for s in sections] != list(range(2, 9)):
        raise ReportError('sections的id必须按2、3、4、5、6、7、8排序且不重复')
    for s in sections:
        loc = f'section[{s["id"]}]'
        if s.get('status') not in STATUSES:
            raise ReportError(f'{loc}.status不是有效核查状态')
        for key in ('location', 'observation'):
            required_text(s, key, loc)
        optional_texts(s, ('recommendation', 'sample'), loc)
    extras = data.get('extras', [])
    if not isinstance(extras, list):
        raise ReportError('extras须为数组')
    for i, extra in enumerate(extras):
        if not isinstance(extra, dict):
            raise ReportError('extras每项须为对象')
        for key in ('title', 'location', 'issue', 'recommendation'):
            required_text(extra, key, f'extra[{i}]')
        optional_texts(extra, ('sample',), f'extra[{i}]')
    evidence = data.get('verification')
    if not isinstance(evidence, dict):
        raise ReportError('verification须为对象')
    for key in ('coverage', 'limitations', 'pending'):
        required_text(evidence, key, 'verification')
    sources = evidence.get('sources', [])
    if not isinstance(sources, list):
        raise ReportError('verification.sources须为数组')
    for i, source in enumerate(sources):
        if not isinstance(source, dict):
            raise ReportError('法源须为对象')
        for key in ('title', 'url', 'checked_on'):
            required_text(source, key, f'source[{i}]')
        url = urlparse(source['url'])
        if url.scheme not in ('http', 'https') or not url.netloc or url.username or url.password:
            raise ReportError('法源url须为无凭据的HTTP(S)地址')
        try:
            date.fromisoformat(source['checked_on'])
        except ValueError as exc:
            raise ReportError('法源checked_on须为日期') from exc
    inputs = data.get('input_files', [])
    if not isinstance(inputs, list) or any(not isinstance(p, str) or not p.strip() for p in inputs):
        raise ReportError('input_files须为路径字符串数组')
    validate_issues(data)
    if 'stance' in data:
        stance = data['stance']
        if not isinstance(stance, dict) or stance.get('party') not in ('甲方', '乙方', '我方'):
            raise ReportError('stance.party须为甲方、乙方或我方')
        catalog = load_catalog(stance.get('role'))
        if 'categories' in stance:
            categories = stance['categories']
            available = {m['category'] for m in catalog['modules']}
            if not isinstance(categories, list) or any(not isinstance(c, str) or c not in available for c in categories):
                raise ReportError('stance.categories须为本方实际存在的类别数组')
    for item in report_items(data):
        fixed_references(item, data)
    return data


def fixed_blocks(item: dict, data: dict) -> list[tuple[str, str, str]]:
    result = []
    for match, module in fixed_references(item, data):
        party, role = data['stance']['party'], data['stance']['role']
        usual_party = '甲方' if role == 'buyer' else '乙方'
        prefix = f'{party}固定条款重点' if party == usual_party else f'{party}立场·{"采购方" if role == "buyer" else "销售方"}固定条款重点'
        sources = [code for code in module.get('source_ids', []) if code != module['id']]
        reference = module['id'] + ('／' + '、'.join(sources) if sources else '')
        result.append(('p', f'{prefix}〔{reference}；{match["status"]}〕', match['analysis']))
    return result


def issue_blocks(data: dict, section: int | str) -> list[tuple[str, str, str]]:
    result = []
    for issue in data.get('issues', []):
        if issue['section'] != section:
            continue
        result.append(('h2', '', f"{issue['id']}｜{issue['title']}〔{ISSUE_TYPES[issue['type']]}〕"))
        result.append(('p', '原稿位置', issue['location']))
        if issue.get('original_excerpt'):
            result.append(('p', '原文', issue['original_excerpt']))
        result.append(('p', '详细说明', issue['detail']))
        result.extend(fixed_blocks(issue, data))
        label = '文字修订建议（未直接改原稿）' if issue['type'] == 'revise' else '注意事项／处理建议'
        result.append(('p', label, issue['suggestion']))
    return result


def blocks(data: dict) -> list[tuple[str, str, str]]:
    """One semantic sequence drives both formats: kind, label, text."""
    stance = data.get('stance')
    perspective = (f"{stance['party']}（{'采购方' if stance['role'] == 'buyer' else '销售方'}）" if stance else data['perspective'])
    result = [('title', '', data['title']), ('meta', '', f"{data['review_date']}　审核立场：{perspective}")]
    result.append(('h1', '', '1、基本情况'))
    for key, label in [('product', '产品'), ('quantity', '数量'), ('price', '价格'), ('delivery', '交付时间'), ('conclusion', '判断')]:
        if data['basics'].get(key):
            result.append(('p', label, data['basics'][key]))
    result.extend(fixed_blocks(data['basics'], data))
    result.extend(issue_blocks(data, 1))
    for section in data['sections']:
        result.append(('h1', '', f"{section['id']}、{HEADINGS[section['id']]}"))
        result.append(('p', f"现状〔{section['status']}；{section['location']}〕", section['observation']))
        if section.get('recommendation'):
            result.append(('p', '意见', section['recommendation']))
        result.extend(fixed_blocks(section, data))
        if section.get('sample'):
            result.append(('p', '备注／建议写法（未确认）', section['sample']))
        result.extend(issue_blocks(data, section['id']))
    if data.get('extras') or any(i['section'] == 'extra' for i in data.get('issues', [])):
        result.append(('h1', '', '框架外明显经济及法律问题'))
        for extra in data['extras']:
            result += [('h2', '', extra['title']), ('p', f"问题〔{extra['location']}〕", extra['issue']), ('p', '意见', extra['recommendation'])]
            result.extend(fixed_blocks(extra, data))
            if extra.get('sample'):
                result.append(('p', '备注／建议写法（未确认）', extra['sample']))
        result.extend(issue_blocks(data, 'extra'))
    result.append(('h1', '', '核验与待确认'))
    for key, label in [('coverage', '已核范围'), ('limitations', '范围限制'), ('pending', '待确认')]:
        result.append(('note', label, data['verification'][key]))
    result.append(('note', '', '本报告仅供审阅；建议写法不代表双方已达成合意，不改变原合同。'))
    return result


def markdown(data: dict) -> str:
    content = []
    for kind, label, text in blocks(data):
        prefix = {'title': '# ', 'h1': '## ', 'h2': '### '}.get(kind, '')
        content.append(prefix + (f'**{label}：** ' if label else '') + text)
    if data['verification'].get('sources'):
        content.append('**核验依据：**')
        content.extend(f"- [{s['title']}]({s['url']})（核验日：{s['checked_on']}）" for s in data['verification']['sources'])
    return '\n\n'.join(content) + '\n'


def word(data: dict) -> bytes:
    try:
        from docx import Document
        from docx.shared import Cm, Pt, RGBColor
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.text import WD_ALIGN_PARAGRAPH
    except ImportError as exc:
        raise ReportError('缺少python-docx；请使用宿主bundled Python或既有hetong Office隔离环境，不自动安装依赖') from exc
    doc = Document()
    for style in doc.styles:
        for border in list(style.element.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(1.8)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.footer_distance = Cm(.8)

    def fonts(style, name, size, bold=False):
        style.font.name = name
        style.font.size = Pt(size)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rf = style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            rf.set(qn('w:' + attr), name)
        for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
            rf.attrib.pop(qn('w:' + attr), None)

    fonts(doc.styles['Normal'], '宋体', 11)
    p = doc.styles['Normal'].paragraph_format
    p.line_spacing, p.space_after, p.widow_control = 1.25, Pt(5), True
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for name, size, before, after in [('Title', 18, 0, 10), ('Heading 1', 13, 10, 5), ('Heading 2', 11, 6, 4)]:
        fonts(doc.styles[name], '黑体', size, True)
        f = doc.styles[name].paragraph_format
        f.keep_with_next, f.space_before, f.space_after, f.line_spacing = True, Pt(before), Pt(after), 1.1
    fonts(doc.styles['Caption'], '宋体', 9.5)
    doc.styles['Caption'].font.italic = False
    doc.styles['Caption'].paragraph_format.line_spacing = 1.2
    doc.styles['Caption'].paragraph_format.keep_with_next = False
    for kind, label, text in blocks(data):
        style = {'title': 'Title', 'h1': 'Heading 1', 'h2': 'Heading 2', 'meta': 'Caption', 'note': 'Caption'}.get(kind, 'Normal')
        para = doc.add_paragraph(style=style)
        if label:
            para.add_run(label + '：').bold = True
        para.add_run(text)
    for source in data['verification'].get('sources', []):
        p = doc.add_paragraph(style='Caption')
        link = OxmlElement('w:hyperlink')
        link.set(qn('r:id'), p.part.relate_to(source['url'], 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
        run = OxmlElement('w:r')
        props = OxmlElement('w:rPr')
        rf = OxmlElement('w:rFonts')
        for attr in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            rf.set(qn('w:' + attr), '宋体')
        props.append(rf)
        run.append(props)
        text = OxmlElement('w:t')
        text.text = source['title']
        run.append(text); link.append(run); p._p.append(link)
        p.add_run(f"（核验日：{source['checked_on']}）")
    fonts(doc.styles['Footer'], '宋体', 9.5)
    p = sec.footer.paragraphs[0]
    p.style = doc.styles['Footer']; p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run('第 ')
    for field_name, suffix in [('PAGE', ' 页／共 '), ('NUMPAGES', ' 页')]:
        field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), field_name)
        p._p.append(field); p.add_run(suffix)
    doc.core_properties.title = data['title']
    doc.core_properties.subject = 'XQY八项快速审核'
    doc.core_properties.author = doc.core_properties.last_modified_by = ''
    output = io.BytesIO(); doc.save(output)
    return output.getvalue()


def output_paths(source: Path, output: Path, data: dict) -> list[Path]:
    if output.suffix.lower() != '.docx':
        raise ReportError('--out必须为.docx')
    paths = [output, output.with_suffix('.md'), output.with_suffix('.build.json')]
    protected = {source.resolve()}
    for value in data.get('input_files', []):
        item = Path(value).expanduser()
        protected.add((item if item.is_absolute() else source.parent / item).resolve())
    for path in paths:
        if '原件' in path.parts or '原件' in path.resolve().parts or path.resolve() in protected:
            raise ReportError(f'禁止覆盖输入或写入原件目录：{path}')
        if path.exists() or path.is_symlink():
            raise ReportError(f'输出已存在，请使用新版本文件名：{path}')
    return paths


def build(source: Path, output: Path, dry_run=False) -> dict:
    raw = source.read_bytes()
    data = validate(json.loads(raw.decode('utf-8-sig')))
    paths = output_paths(source, output, data)
    result = {'schema': SCHEMA, 'dry_run': dry_run, 'outputs': [str(p) for p in paths], 'input_sha256': hashlib.sha256(raw).hexdigest(), 'sections': 8, 'body_font': '宋体', 'body_size_pt': 11, 'substantive_review_verified': False, 'visual_review_verified': False}
    if 'stance' in data:
        result['stance'] = data['stance']
        role = data['stance']['role']
        catalog = load_catalog(role)
        result['fixed_clause_catalog'] = {'id': catalog['catalog_id'], 'role': role, 'sha256': hashlib.sha256((CATALOG_DIR / f'{role}-clause-catalog.json').read_bytes()).hexdigest()}
        result['fixed_clause_references'] = [{**module, 'status': match['status']} for item in report_items(data) for match, module in fixed_references(item, data)]
    if dry_run:
        return result
    payloads = [word(data), markdown(data).encode('utf-8')]
    result['output_sha256'] = {str(p): hashlib.sha256(blob).hexdigest() for p, blob in zip(paths, payloads)}
    payloads.append((json.dumps(result, ensure_ascii=False, indent=2) + '\n').encode('utf-8'))
    created = []
    try:
        for path, blob in zip(paths, payloads):
            path.parent.mkdir(parents=True, exist_ok=True)
            with path.open('xb') as stream:
                created.append(path)
                stream.write(blob)
    except BaseException:
        for path in created:
            path.unlink(missing_ok=True)
        raise
    return result


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, required=True)
    parser.add_argument('--out', type=Path, required=True)
    parser.add_argument('--dry-run', action='store_true')
    args = parser.parse_args()
    try:
        result = build(args.input.expanduser().absolute(), args.out.expanduser().absolute(), args.dry_run)
        print(json.dumps(result, ensure_ascii=False, indent=2))
        return 0
    except (ReportError, OSError, json.JSONDecodeError, UnicodeError) as exc:
        print(f'错误：{exc}', file=sys.stderr)
        return 2


if __name__ == '__main__':
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding='utf-8')
        except (AttributeError, OSError):
            pass
    raise SystemExit(main())
