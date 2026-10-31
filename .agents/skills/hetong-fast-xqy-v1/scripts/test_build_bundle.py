#!/usr/bin/env python3
"""End-to-end fictional report/comment pairing and source preservation tests."""
from copy import deepcopy
import importlib.util
import io
import json
from pathlib import Path
import tempfile
import unittest
from zipfile import ZipFile
import xml.etree.ElementTree as ET

from docx import Document

HERE = Path(__file__).resolve().parent
spec = importlib.util.spec_from_file_location('xqy_bundle_tests', HERE / 'build_bundle.py')
bundle = importlib.util.module_from_spec(spec); spec.loader.exec_module(bundle)
NS = bundle.NS


def make_case(root):
    root = Path(root); root.mkdir(parents=True, exist_ok=True)
    source = root / '虚构原稿.docx'
    doc = Document(); doc.add_paragraph('虚构设备采购合同')
    doc.add_paragraph('乙方收到货款后30日内开具发票。')
    doc.add_paragraph('双方指定经办人员确认序列号清单。')
    doc.add_paragraph('甲方应在到货后当场提出全部异议，否则视为全面合格。')
    doc.add_paragraph('本段仅供历史审阅保留测试。')
    doc.save(source)
    data = json.loads((HERE.parent / 'assets/example-report.json').read_text())
    data['title'] = '虚构采购合同 双文件审阅样张'
    data['basics']['conclusion'] = '纯虚构样张；通过编号联系报告详细建议和原稿批注。'
    for item in bundle.report.report_items(data):
        item.pop('fixed_clauses', None); item.pop('sample', None)
    for item in data['sections']:
        if item['id'] not in (2, 4, 5):
            item.update(status='待核实', location='虚构样张范围之外',
                        observation='本样张仅展示三项编号对应，未提供本主题的完整背景。',
                        recommendation='真实审核须按完整合同核对；本样张不就此作实体结论。')
    data['sections'][0]['observation']='原稿约定收款后30日开票，但缺少错误更正与逾期处理。'
    data['sections'][2].update(status='已具备', observation='原稿已有经办人员确认序列号清单的要求，正文保留。')
    data['sections'][3]['observation']='原稿将未当场提出全部异议视为全面合格。'
    data['extras'] = []
    data['bundle'] = {'source': source.name, 'source_sha256': bundle.sha(source.read_bytes()), 'author': '合同审阅',
                      'report_name': '虚构合同_v1_XQY审核报告.docx', 'comments_name': '虚构合同_v1_XQY批注稿.docx'}
    data['issues'] = [
        {'id':'R-001','section':2,'type':'revise','title':'补齐开票及逾期责任','location':'开票条款',
         'original_excerpt':'乙方收到货款后30日内开具发票。',
         'detail':'原稿仅列开票时间，未说明逾期或不合格发票的处理。以下内容是拟议方案，参数尚未确认。',
         'suggestion':'建议替换为：乙方应在收到对应货款后【待确认：期限】内交付合法有效发票；逾期按【待确认：基数及比例】承担责任，错误发票应及时更正。',
         'comment_summary':'开票条款缺少逾期及更正机制，建议按报告文本完善。',
         'anchors':[{'para':'RL0002','text':'乙方收到货款后30日内开具发票。'}]},
        {'id':'R-002','section':5,'type':'revise','title':'区分外观与性能验收','location':'验收条款',
         'detail':'到货当场难以发现全部内部缺陷，原稿将未当场异议等同全面合格，影响后续质量救济。',
         'suggestion':'建议删除“否则视为全面合格”，改为：签收仅确认数量及可见外观，性能验收按双方书面确认的方案实施；隐蔽缺陷及质保权利不因签收受影响。',
         'comment_summary':'当场未异议不应直接视为性能全面合格，建议区分两类验收。',
         'anchors':[{'para':'RL0004','text':'否则视为全面合格'}]},
        {'id':'R-003','section':4,'type':'attention','title':'交接清单的人员及证据','location':'资料交接条款',
         'detail':'原稿已有序列号清单确认义务，可予保留；履行时仍需核实经办人员授权并留存清单和签收记录。',
         'suggestion':'正文保留。交接前核验双方经办人员授权，按批次保存序列号清单及确认记录。',
         'comment_summary':'本处正文保留；实际交接时核验经办授权并留存按批次签收证据。',
         'anchors':[{'para':'RL0003','text':'双方指定经办人员确认序列号清单。'}]}
    ]
    # Initial IDs follow report order; later versions retain existing IDs.
    data['issues'][1]['id'] = 'R-003'
    data['issues'][2]['id'] = 'R-002'
    data['issues'] = [data['issues'][0], data['issues'][2], data['issues'][1]]
    path = root / 'review.json'; path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')
    return path, source, data


def save_case(path, data):
    path.write_text(json.dumps(data, ensure_ascii=False), encoding='utf-8')


class BundleTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(); self.root = Path(self.tmp.name)
        self.input, self.source, self.data = make_case(self.root)
        self.out = self.root / '交付_v1'

    def tearDown(self): self.tmp.cleanup()

    def test_report_native_comments_and_body_preserved(self):
        source_bytes = self.source.read_bytes()
        result = bundle.build(self.input, self.out)
        self.assertEqual(self.source.read_bytes(), source_bytes)
        self.assertEqual(result['comment_count'], 3)
        self.assertEqual(len(list(self.out.glob('*.docx'))), 2)
        with ZipFile(self.out / result['comments']) as z:
            comments = ET.fromstring(z.read('word/comments.xml'))
            text = ''.join(comments.itertext())
            for issue in self.data['issues']:
                self.assertIn(issue['id'], text)
                self.assertIn(issue['comment_summary'], text)
            self.assertIn(result['report'], text)
            root = ET.fromstring(z.read('word/document.xml'))
            self.assertEqual(len(root.findall('.//w:ins', NS)), 0)
            self.assertEqual(len(root.findall('.//w:del', NS)), 0)
        with ZipFile(self.out / result['report']) as z:
            text = ''.join(ET.fromstring(z.read('word/document.xml')).itertext())
            for issue in self.data['issues']:
                self.assertIn(issue['suggestion'], text)
        self.assertTrue(bundle.verify_bundle(self.out)['ok'])

    def test_dry_run_no_output(self):
        bundle.build(self.input, self.out, True)
        self.assertFalse(self.out.exists())

    def test_duplicate_id_or_missing_suggestion_rejected(self):
        for key, value in [('id', 'R-001'), ('suggestion', '')]:
            data = deepcopy(self.data); data['issues'][1][key] = value; save_case(self.input, data)
            with self.assertRaises(bundle.report.ReportError): bundle.build(self.input, self.out)
            self.assertFalse(self.out.exists())

    def test_wrong_anchor_or_missing_anchor_rejected(self):
        for anchors in ([{'para':'RL9999','text':'不存在'}], []):
            data = deepcopy(self.data); data['issues'][0]['anchors'] = anchors; save_case(self.input, data)
            with self.assertRaises(bundle.report.ReportError): bundle.build(self.input, self.out)
            self.assertFalse(self.out.exists())

    def test_source_drift_rejected(self):
        self.source.write_bytes(self.source.read_bytes() + b'drift')
        with self.assertRaises(bundle.report.ReportError): bundle.build(self.input, self.out)
        self.assertFalse(self.out.exists())

    def test_two_issues_same_paragraph_grouped_with_both_references(self):
        issue = deepcopy(self.data['issues'][0]); issue['id']='R-004'; issue['type']='attention'
        issue['anchors']=[{'para':'RL0002','text':'收到货款'}]
        issue['comment_summary']='留存到款证明并确认开票起算。'
        self.data['issues'].append(issue); save_case(self.input,self.data)
        result=bundle.build(self.input,self.out)
        self.assertEqual(result['comment_count'],3)
        self.assertEqual(result['mappings'][0]['issue_ids'],['R-001','R-004'])
        self.assertTrue(bundle.verify_bundle(self.out)['ok'])

    def test_one_issue_multiple_anchors_keeps_one_report_id(self):
        self.data['issues'][0]['anchors'].append({'para':'RL0003','text':'序列号清单'})
        self.data['issues']=self.data['issues'][:1];save_case(self.input,self.data)
        result=bundle.build(self.input,self.out)
        self.assertEqual(result['issue_ids'],['R-001'])
        self.assertEqual(result['comment_count'],2)

    def test_existing_output_rejected(self):
        self.out.mkdir(); (self.out/'用户文件.txt').write_text('保留')
        with self.assertRaises(bundle.report.ReportError): bundle.build(self.input,self.out)
        self.assertEqual((self.out/'用户文件.txt').read_text(),'保留')

    def test_tampered_report_rejected(self):
        result=bundle.build(self.input,self.out)
        report=self.out/result['report']; report.write_bytes(report.read_bytes()+b'changed')
        with self.assertRaises(bundle.report.ReportError): bundle.verify_bundle(self.out)

    def test_empty_issue_copy_is_explicit_and_exact(self):
        self.data['issues']=[];save_case(self.input,self.data)
        result=bundle.build(self.input,self.out)
        self.assertEqual(result['comment_count'],0)
        self.assertEqual((self.out/result['comments']).read_bytes(),self.source.read_bytes())
        self.assertTrue(bundle.verify_bundle(self.out)['ok'])

    def test_preserves_prior_comments_and_revisions(self):
        engine=bundle.engine;ctx=engine.load_context(str(self.source),for_write=True)
        oldplan=engine.build_plan(ctx,[{'id':'old','para':'RL0005','old':'本段仅供历史审阅保留测试。','comment':'历史批注不得删除'}])
        oldbytes,_=engine.build_result(ctx,oldplan,'旧审阅',timestamp='2000-01-01T00:00:00Z')
        with ZipFile(io.BytesIO(oldbytes)) as z:
            members={name:z.read(name) for name in z.namelist()}
        doc=ET.fromstring(members['word/document.xml']);p=doc.findall('.//w:p',NS)[-1]
        ins=ET.SubElement(p,'{'+NS['w']+'}ins',{'{'+NS['w']+'}id':'42','{'+NS['w']+'}author':'旧修订','{'+NS['w']+'}date':'2000-01-01T00:00:00Z'})
        ET.SubElement(ET.SubElement(ins,'{'+NS['w']+'}r'),'{'+NS['w']+'}t').text='历史新增内容'
        members['word/document.xml']=ET.tostring(doc,encoding='utf-8',xml_declaration=True)
        with ZipFile(self.source,'w') as z:
            for name,blob in members.items():z.writestr(name,blob)
        self.data['bundle']['source_sha256']=bundle.sha(self.source.read_bytes());save_case(self.input,self.data)
        result=bundle.build(self.input,self.out)
        with ZipFile(self.out/result['comments']) as z:
            text=''.join(ET.fromstring(z.read('word/comments.xml')).itertext())
            self.assertIn('历史批注不得删除',text)
            doc=ET.fromstring(z.read('word/document.xml'))
            oldins=doc.find(".//w:ins[@w:id='42']",NS)
            self.assertIsNotNone(oldins);self.assertEqual(''.join(oldins.itertext()),'历史新增内容')
        self.assertTrue(bundle.verify_bundle(self.out)['ok'])


if __name__=='__main__': unittest.main(verbosity=2)
